# -*- coding: utf-8 -*-
"""db.json → pikogen_db.xlsx（Googleスプレッドシートに読み込ませるDB雛形）"""
import json, csv, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

D = json.load(open('tools/db.json', encoding='utf-8'))
os.makedirs('tools/csv', exist_ok=True)

INK = "14140F"; YEL = "F2D000"; HEAD = "1F1F19"; ALT = "FBF8E8"

def sheet(wb, name, headers, rows, widths=None, note=None, freeze="A2"):
    ws = wb.create_sheet(name)
    r0 = 1
    if note:
        ws.cell(row=1, column=1, value=note).font = Font(size=9, italic=True, color="77775F")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(headers)))
        r0 = 2
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=r0, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=HEAD)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = Border(bottom=Side(style="thin", color=INK))
    ws.row_dimensions[r0].height = 26
    for ri, row in enumerate(rows, r0 + 1):
        for ci, v in enumerate(row, 1):
            if isinstance(v, (dict, list)):
                v = json.dumps(v, ensure_ascii=False)
            c = ws.cell(row=ri, column=ci, value=v)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.font = Font(size=10)
            if (ri - r0) % 2 == 0:
                c.fill = PatternFill("solid", fgColor=ALT)
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=r0 + 1, column=1)
    ws.sheet_view.showGridLines = False
    # CSV も書き出す（検証用）
    with open(f'tools/csv/{name}.csv', 'w', encoding='utf-8', newline='') as f:
        w = csv.writer(f)
        w.writerow(headers)
        for row in rows:
            w.writerow([json.dumps(x, ensure_ascii=False) if isinstance(x, (dict, list)) else ('' if x is None else x) for x in row])
    return ws

wb = Workbook(); wb.remove(wb.active)
W = D['weights']

# ---------- 設定 ----------
SETTINGS = [
    ('version','1.0.0','DBの版。変更したら上げる'),
    ('ip',W['ip'],'スコア重み：IP一致'),
    ('verb',W['verb'],'スコア重み：動詞一致'),
    ('emo',W['emo'],'スコア重み：感情一致'),
    ('aud',W['aud'],'スコア重み：対象一致'),
    ('tech',W['tech'],'スコア重み：技術適合'),
    ('warningRatio',0.8,'整合度がこれを下回ったら警告'),
    ('maxTwist',1,'1ゲームに入れる刺激の数（変えない）'),
    ('bpm',D['picoSound']['bpm'],'ピコ太郎BGMのBPM'),
    ('picoSeCount',len(D['picoSe']),'SEの数'),
]
sheet(wb,'設定',['key','value','説明'],[list(x) for x in SETTINGS],[20,14,46],'DB全体の設定値。')

# ---------- マスタ ----------
sheet(wb, 'マスタ_対象', ['id','label','note','dev'],
      [[a['id'], a['label'], a['note'], a.get('dev','both')] for a in D['audiences']], [20,22,52,10],
      'STEP「誰に」の選択肢。dev は pc / tap / both。')

sheet(wb, 'マスタ_IP', ['id','group','label','mech','note'],
      [[v['id'], v['g'], v['label'], v['mech'], v['note']] for v in D['ips']], [16,16,30,8,54],
      'mech：◎=仕組みにできる／○=構造に効く／△=見た目のみ。ここがIP強度判定の根拠。')

sheet(wb, 'マスタ_感情', ['id','label','ok','ng','tone'],
      [[e['id'], e['label'], e['ok'], e['ng'], e['tone']] for e in D['emotions']], [16,22,26,26,16],
      '成功時・失敗時にどう感じてほしいか。演出とSEの方向を決める。')

rows = []
for g in D['verbs']:
    for v in g['verbs']:
        rows.append([v['id'], g['g'], v['label'], v['pat'], v['d']])
sheet(wb, 'マスタ_動詞', ['id','group','label','pat','difficulty'], rows, [14,20,20,10,10],
      'pat は既定パターンID。difficulty は 1〜3。')

sheet(wb, 'マスタ_型', ['id','label','engines','level'],
      [[t['id'], t['label'], ','.join(t['engines']), t['level']] for t in D['types']], [18,34,22,8],
      'engines はカンマ区切り。ここに一致するエンジンの案だけが出る。')

sheet(wb, 'マスタ_欲求', ['id','label','note'],
      [[w['id'], w['label'], w['note']] for w in D['wantTo']], [18,34,34], '「○○したい人」の選択肢。')

sheet(wb, 'マスタ_認知', ['id','label','tags'],
      [[p['id'], p['label'], ','.join(p['tags'])] for p in D['perception']], [18,52,44],
      '「ピコ太郎は○○なIP」の選択肢。tags は IP価値ID。')

sheet(wb, 'マスタ_なぜ', ['id','short','text','tags'],
      [[w['id'], w['short'], w['text'], ','.join(w['tags'])] for w in D['whyPico']], [18,26,74,40],
      '第3回で出てくる「なぜピコ太郎なのか」の候補文。学生はここから選ぶ。')

# ---------- パターン ----------
ph = ['id','name','family','engine','dim','level','input','duration','difficulty','oneline','action',
      'success_tpl','fail_tpl','loop','retry_reason','ip_hook','verb_group','target','life','fall','interval',
      'speedup','image_comp','image_mood','code_notes','bgm','active']
prow = []
for p in D['patterns']:
    n = p['nums']
    prow.append([p['id'], p['name'], p['family'], p['engine'], p.get('dim','2D'), p['level'], p['input'],
                 p['duration'], p['difficulty'], p['oneline'], p['action'], p['successTpl'], p['failTpl'],
                 p['loop'], p['retry'], p['ipHook'], p['verbGroup'], n['target'], n['life'], n['fall'],
                 n['interval'], n['speedup'], p['imageComp'], p['imageMood'], p['codeNotes'], p['bgm'], 'TRUE'])
sheet(wb, 'パターン', ph, prow, [8,24,12,8,6,6,16,8,8,44,54,30,26,34,36,50,16,7,6,7,8,8,60,44,60,34,7],
      '1行=1ゲーム案。文章はここを直す。active=FALSE で非表示。')

c_ip = [[p['id'], k, v] for p in D['patterns'] for k, v in p['ip'].items()]
sheet(wb, 'パターン_IP', ['pattern_id','ip_id','weight'], c_ip, [10,18,9], '重みは 0〜1。大きいほど「そのIP価値を強く使っている」判定。')

c_ui = [[p['id'], u[0], u[1]] for p in D['patterns'] for u in p['ui']]
sheet(wb, 'パターン_UI', ['pattern_id','ui_id','req'], c_ui, [10,16,8], 'req は 必須／推奨／任意。')

c_as = [[p['id'], a] for p in D['patterns'] for a in p['assets']]
sheet(wb, 'パターン_素材', ['pattern_id','asset_id'], c_as, [10,18], '素材マスタのIDを参照。')

c_se = [[p['id'], s] for p in D['patterns'] for s in p['se']]
sheet(wb, 'パターン_音', ['pattern_id','se_id'], c_se, [10,18], 'SEマスタのIDを参照。')

c_em = [[p['id'], e] for p in D['patterns'] for e in p['emo']]
sheet(wb, 'パターン_感情', ['pattern_id','emo_id'], c_em, [10,18], 'この案が刺さる感情。')

c_au = [[p['id'], a] for p in D['patterns'] for a in p['aud']]
sheet(wb, 'パターン_対象', ['pattern_id','aud_id'], c_au, [10,18], 'この案が向く対象。')

# ---------- エンジン ----------
erows = []
for eid, e in D['engines'].items():
    erows.append([eid, e['name'], e['short'], e['scene'], json.dumps(e['comps'], ensure_ascii=False),
                  e['flow'], json.dumps(e['tuning'], ensure_ascii=False), ','.join(e.get('juice', []))])
sheet(wb, 'エンジン', ['id','name','short','scene','comps','flow','tuning','juice'], erows,
      [8,26,40,60,50,52,60,34], '実装テンプレート。sceneは改行保持。comps/tuningはJSON配列。')

grow = []
for eid, e in D['engines'].items():
    for i, s in enumerate(e['mid'], 1):
        grow.append([eid, i, s['h'], s['w'], s['d'], s['c'], s['t'], s['m'], s['p']])
sheet(wb, 'エンジン_手引き', ['engine_id','order','heading','work','done','check','trap','minutes','prefab'],
      grow, [9,6,26,64,34,34,54,8,18], '制作手引きの工程。order順に表示される。')

jrow = [[k, v['name'], v['when'], v['how'], v['tune']] for k, v in D['juice'].items()]
sheet(wb, 'ジュース', ['id','name','when','how','tune'], jrow, [16,20,34,64,40], '演出（ジュース）の定義。')

# ---------- 素材・UI・音 ----------
arow = [[k, v['role'], v['name'], v['size'], v['fake'], v['pri'], v['path']] for k, v in D['assets'].items()]
sheet(wb, '素材', ['id','role','name','size','fake','priority','path'], arow, [16,12,26,16,56,8,28],
      '仮素材の作り方まで書く。priority は 低／中／高。')

urow = [[k, v['label'], v['pos'], v['why']] for k, v in D['ui'].items()]
sheet(wb, 'UI', ['id','label','pos','why'], urow, [12,16,16,52], 'UIの表示位置と役割。')

srow = [[k, v['label'], v.get('file',''), v.get('len',''), v['when'], v['tone'], v.get('vol','')] for k, v in D['se'].items()]
sheet(wb, 'SE', ['id','label','file','length','when','tone','volume'], srow, [14,16,18,10,30,36,8],
      'ピコ太郎らしさはSEで決まる。Pen/Apple/Pineappleは耳で区別できる音にする。')

brow = [[k, v['bpm'], v['key'], v['mood'], v['instr'], v['loop']] for k, v in D['bgm'].items()]
sheet(wb, 'BGM', ['engine_id','bpm','key','mood','instr','loop_sec'], brow, [10,8,8,26,50,10],
      'BGMの基準値。エンジンごとに変える。')

prow2 = [[k, v['name'], v['fn'], v['cat'], v['level'], ','.join(v['deps']), v['done'], v['check'], v['trap'], v['min']]
         for k, v in D['prefabs'].items()]
sheet(wb, 'Prefab', ['id','name','fn','cat','level','deps','done','check','trap','minutes'], prow2,
      [20,20,22,10,6,22,40,40,54,8], '機能名を必ず併記する。levelは1〜3。')

trow = [[t['id'], t['type'], t['name'], t['one'], t['what'], t['how'], t['effect'], t['tune'], ','.join(t['fit']), ','.join(t['verbs'])]
        for t in D['twists']]
sheet(wb, 'ツイスト', ['id','type','name','one','what','how','effect','tune','fit_engines','fit_verbs'],
      trow, [14,10,20,30,50,50,40,36,26,34], '刺激の1手。1ゲームに1つだけ入れる。type: obstacle/limit/dual/reverse/tease')

sheet(wb, 'ツイスト_動詞', ['verb','twist_id'],
      [[k, v] for k, v in D['subverbTwist'].items()], [20,16], '副動詞から刺激を選ぶときの優先表。')

frow = [[f['id'], f['y'], f['title'], f['fact'], f['insight'], f['src'], ','.join(f['tags'])] for f in D['facts']]
sheet(wb, 'IP事実', ['id','year','title','fact','insight','source','tags'], frow,
      [14,8,30,70,60,54,34], '提案書の「共感」で引用する事実。出典URLは必ず残す。')

# 並び順：設定 → マスタ → パターン → エンジン → 素材/音 → その他
order = ['設定','マスタ_なぜ','マスタ_欲求','マスタ_認知','マスタ_対象','マスタ_IP','マスタ_感情','マスタ_動詞','マスタ_型',
         'パターン','パターン_IP','パターン_UI','パターン_素材','パターン_音','パターン_感情','パターン_対象',
         'エンジン','エンジン_手引き','ツイスト','ツイスト_動詞',
         '素材','UI','SE','BGM','Prefab','ジュース','IP事実']
wb._sheets = [wb[n] for n in order]
wb.active = 0
wb.save('pikogen_db.xlsx')
print('sheets:', len(wb.sheetnames))
print('\n'.join(wb.sheetnames))
